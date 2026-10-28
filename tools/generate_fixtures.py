"""Generate the synthetic test fixtures under tests/fixtures.

Run from the repo root with the dev extras installed:

    .venv/Scripts/python tools/generate_fixtures.py

The data is invented. Property names, amounts and IDs do not correspond to
any real journal entry. The clean sample is written as CSV, XLSX, XLS and ODS
so the readers can be checked for identical output, plus one XLSX with a
second sheet to exercise the extra_sheets flag.
"""

from __future__ import annotations

import csv
import datetime as dt
from pathlib import Path

HEADER = [
    "subsidiary",
    "postingperiod",
    "journalItemheader_memo",
    "tranDate",
    "journalItemLine_accountRef",
    "journalItemLine_debitAmount",
    "journalItemLine_creditAmount",
    "journalItemLine_memo",
    "journalItemLine_departmentRef",
    "journalItemLine_classRef",
    "journalItemLine_locationRef",
]

PP = dt.date(2026, 10, 15)
TD = dt.date(2026, 10, 20)

Row = list[object]


def line(
    sub: str,
    memo: str,
    acct: object,
    debit: float | None,
    credit: float | None,
    dept: str,
    loc: str,
    *,
    lmemo: str = "",
    cls: str = "",
) -> Row:
    return [int(sub), PP, memo, TD, acct, debit, credit, lmemo, dept, cls, int(loc)]


# 1. Clean hotel entry, path-style departments, balanced.
CLEAN_MEMO = "100001.200001.4000001"
CLEAN: list[Row] = [
    line("41", CLEAN_MEMO, 279, None, 1500.00, "Administration:G&A:General Manager", "61"),
    line("41", CLEAN_MEMO, 279, None, 4250.75, "Operations:Hotel Ops:Front Office:Front Desk Agent", "61"),
    line("41", CLEAN_MEMO, 279, None, 6100.25, "Operations:Hotel Ops:Housekeeping:Housekeeper", "61"),
    line("41", CLEAN_MEMO, 761, 12.50, None, "Operations:Hotel Ops:Front Office:Front Desk Agent", "61"),
    line("41", CLEAN_MEMO, 767, 1300.00, None, "Administration:G&A:General Manager", "61"),
    line("41", CLEAN_MEMO, 767, 3700.00, None, "Operations:Hotel Ops:Front Office:Front Desk Agent", "61"),
    line("41", CLEAN_MEMO, 767, 5400.00, None, "Operations:Hotel Ops:Housekeeping:Housekeeper", "61"),
    line("41", CLEAN_MEMO, 773, 200.00, None, "Administration:G&A:General Manager", "61"),
    line("41", CLEAN_MEMO, 773, 538.25, None, "Operations:Hotel Ops:Front Office:Front Desk Agent", "61"),
    line("41", CLEAN_MEMO, 773, 700.25, None, "Operations:Hotel Ops:Housekeeping:Housekeeper", "61"),
]
# totals: credit 11851.00, debit 11851.00

# 2. Column shift: every reference column equals the subsidiary.
SHIFT_MEMO = "100002.200002.4000002"
SHIFT: list[Row] = [
    [42, PP, SHIFT_MEMO, TD, 42, None, 1000.00, 42, 42, 42, 42],
    [42, PP, SHIFT_MEMO, TD, 42, None, 2000.00, 42, 42, 42, 42],
    [42, PP, SHIFT_MEMO, TD, 42, 2500.00, None, 42, 42, 42, 42],
    [42, PP, SHIFT_MEMO, TD, 42, 500.00, None, 42, 42, 42, 42],
]

# 3. Multi-subsidiary store group with a text account name.
MULTI_MEMO = "100003.200003.4000003"
MULTI: list[Row] = [
    line("51", MULTI_MEMO, "Payroll Clearing", None, 3000.00, "", "71"),
    line("51", MULTI_MEMO, 767, 2700.00, None, "", "71"),
    line("51", MULTI_MEMO, 773, 300.00, None, "", "71"),
    line("52", MULTI_MEMO, "Payroll Clearing", None, 1200.00, "1", "72"),
    line("52", MULTI_MEMO, 767, 1080.00, None, "1", "72"),
    line("52", MULTI_MEMO, 773, 120.00, None, "1", "72"),
]

# 4. Mixed department styles, off by one cent, memo does not match filename.
MIXED_MEMO = "999999.888888.4000004"
MIXED: list[Row] = [
    line("43", MIXED_MEMO, 279, None, 800.00, "101", "63"),
    line("43", MIXED_MEMO, 279, None, 400.00, "HOUSEKEEPING", "63"),
    line("43", MIXED_MEMO, 767, 700.00, None, "101", "63"),
    line("43", MIXED_MEMO, 767, 350.00, None, "Operations:Hotel Ops:Housekeeping", "63"),
    line("43", MIXED_MEMO, 773, 149.99, None, "101", "63"),
]

# 5. Wrong layout entirely.
WRONG_HEADER = [["Employee", "Hours", "Rate"], ["A. Person", 40, 21.5]]


def to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.date):
        return value.strftime("%m/%d/%Y")
    if isinstance(value, float):
        return f"{value:.2f}"
    return str(value)


def write_csv(path: Path, rows: list[Row], header: list[str] = HEADER) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(header)
        for row in rows:
            writer.writerow([to_text(v) for v in row])


def write_xlsx(path: Path, rows: list[Row], extra_sheet: bool = False) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "JE"
    ws.append(HEADER)
    for row in rows:
        ws.append([v for v in row])
    for cell in ws["B"][1:] + ws["D"][1:]:
        cell.number_format = "MM/DD/YYYY"
    if extra_sheet:
        other = wb.create_sheet("Notes")
        other.append(["This sheet must be ignored."])
    wb.save(path)


def write_xls(path: Path, rows: list[Row]) -> None:
    import xlwt

    wb = xlwt.Workbook()
    ws = wb.add_sheet("JE")
    date_style = xlwt.easyxf(num_format_str="MM/DD/YYYY")
    for c, name in enumerate(HEADER):
        ws.write(0, c, name)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            if value is None:
                continue
            if isinstance(value, dt.date):
                ws.write(r, c, value, date_style)
            else:
                ws.write(r, c, value)
    wb.save(str(path))


def write_ods(path: Path, rows: list[Row]) -> None:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    doc = OpenDocumentSpreadsheet()
    table = Table(name="JE")

    def add_row(values: list[object]) -> None:
        tr = TableRow()
        for value in values:
            if value is None:
                tc = TableCell()
            elif isinstance(value, dt.date):
                tc = TableCell(valuetype="date", datevalue=value.isoformat())
                tc.addElement(P(text=value.strftime("%m/%d/%Y")))
            elif isinstance(value, (int, float)):
                tc = TableCell(valuetype="float", value=str(value))
                tc.addElement(P(text=str(value)))
            else:
                tc = TableCell(valuetype="string")
                tc.addElement(P(text=str(value)))
            tr.addElement(tc)
        table.addElement(tr)

    add_row(HEADER)
    for row in rows:
        add_row(row)
    doc.spreadsheet.addElement(table)
    doc.save(str(path))


def main() -> None:
    out = Path(__file__).resolve().parent.parent / "tests" / "fixtures"
    out.mkdir(parents=True, exist_ok=True)

    write_csv(out / "Sample Hotel JE_100001_200001.CSV", CLEAN)
    write_xlsx(out / "Sample Hotel JE_100001_200001.xlsx", CLEAN)
    write_xlsx(out / "Sample Hotel Two Sheets JE_100001_200001.xlsx", CLEAN, extra_sheet=True)
    write_xls(out / "Sample Hotel JE_100001_200001.xls", CLEAN)
    write_ods(out / "Sample Hotel JE_100001_200001.ods", CLEAN)

    write_csv(out / "Shifted Inn JE_100002_200002.CSV", SHIFT)
    write_csv(out / "Multi Store JE_100003_200003.CSV", MULTI)
    write_csv(out / "Mixed Lodge JE_100004_200004.CSV", MIXED)
    write_csv(out / "random-report.csv", WRONG_HEADER[1:], header=WRONG_HEADER[0])
    (out / "notes.txt").write_text("Not a spreadsheet.\n", encoding="utf-8")

    # A CSV in Windows-1252 with a non-ASCII property name to exercise the decode fallback.
    cp1252 = out / "Café Lodge JE_100005_200005.CSV"
    rows = [
        line("44", "100005.200005.4000005", 279, None, 100.00, "Café", "64"),
        line("44", "100005.200005.4000005", 767, 100.00, None, "Café", "64"),
    ]
    with cp1252.open("w", newline="", encoding="cp1252") as fh:
        writer = csv.writer(fh)
        writer.writerow(HEADER)
        for row in rows:
            writer.writerow([to_text(v) for v in row])

    print(f"wrote fixtures to {out}")
    for p in sorted(out.iterdir()):
        print(f"  {p.name} ({p.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
